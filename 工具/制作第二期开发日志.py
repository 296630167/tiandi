from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '开发日志第二期'
TITLE = '从参加聚光灯比赛到我为天帝镇压世间一切 第二期开发日志'
# 同一稿源同时生成Word与可复制的论坛Markdown。所有图均为本轮实机采集。
PAGES = [
{'heading': None, 'blocks': [
('p', '大家好，我是烤肠。'),
('p', '本作正在参加2026年聚光灯游戏开发比赛，主题是「涌现」。上一期聊了游戏怎么从动物与自然的方向，转到了道纹构筑。这一期，继续和大家说说这段时间做了什么。'),
('section', '本期进展'),
('p', '道纹能拼起来之后，我想继续解决两个问题：玩家能不能看懂自己搭出来的东西，以及这套构筑进入战斗后，能不能形成完整的成长过程。于是，这一版主要围绕道纹通路、战斗地图和操作体验做了一轮调整。'),
('fig', '01_标题', '当前标题页面 游戏仍处于原型开发阶段'),
('p', '游戏的起点也做了个小调整：主角现在只能勉强射出一丝灵力，天然攻击的基础伤害仍然是1。想变强，就得靠天赋、道纹连接和后续成长。名字虽然很长，离“镇压世间一切”也还早，先把眼前这张图打明白。')
]},
{'heading': '一 序章和天赋选择', 'blocks': [
('p', '序章目前使用插画、字幕和转场来讲穿越的故事。底部有进度条，可以调整进度，也可以暂停或跳过，让玩家知道剧情播到了哪里。正式剧情视频还在后续制作范围里。'),
('fig', '02_序章进度', '序章插画与底部进度条'),
('p', '天赋池仍是第一期介绍的十三种，每次随机展示五种。道纹飞近后可以悬停查看效果，确认其中一枚作为画布中心的起点。刷新天赋也保留了，同一轮五个候选不会重复。'),
('fig', '03_源道纹天赋选择', '五选一天赋与悬停详情')
]},
{'heading': '二 让道纹画布更像一张构筑图', 'blocks': [
('p', '道纹界面重新整理成墨青画布、淡纸属性栏和下方藏匣。100×100的六边形格子保留，中心默认解锁，其余每格消耗1点技能点；升级仍然每级给1点。'),
('fig', '06_道纹构筑', '当前道纹画布 接口接通后才会高亮'),
('p', '源道纹的六个接口现在对应独立通路。开局只开放一个，另外五个在10、20、30、40、50级依次解封。力量等基础和普通属性全局计算、同枚去重，攻击形态与五行伤害按各条通路分别计算。有效通路轮流攻击，共用总攻速。'),
('fig', '07_源纹接口解封', '六条通路页签显示尚未解封的接口与等级')
]},
{'heading': '三 放下之前先看看会发生什么', 'blocks': [
('p', '拖动道纹时，右侧会先显示放置后的变化：哪些道纹会亮起，哪些会断开，伤害和攻击形态怎样变化。绿色代表这个格子可以放，真正接通源道纹后才会生效。预览期间不会提前修改构筑。'),
('fig', '25_拖放沙盘对比', '拖放预览 右侧说明当前摆法不会产生有效加成'),
('p', '未接通的道纹也有定位和原因提示。接口需要双向对接，线路只能向同圈或更外圈传导。摆错后可以右键转一格、双击收回，也可以撤销上一步。我希望玩家调整布局时，能看见原因，而不是凭感觉试到它亮。'),
('fig', '08_未接通定位', '定位未接通的道纹 画布与原因提示对应')
]},
{'heading': '四 找道纹和保存方案更方便了', 'blocks': [
('p', '藏匣增加了属性分类、品阶排序和接口筛选。方向可以多选，也可以把“旋转后能匹配”的道纹一起找出来；画布保留回到源点和聚焦已解锁区域，少一点来回翻找。'),
('fig', '09_接口筛选', '属性与接口筛选 可以检查旋转后的匹配'),
('p', '另外加入了三个可以命名的布局方案。记录的是玩家实际拥有的道纹、位置和朝向，载入前会检查库存和格子；洗练后的词条会继续保留。这样可以保存不同想法，之后再回来调整。'),
('fig', '10_布局方案', '三个布局槽位 保存后可检查并载入')
]},
{'heading': '五 属性分类和图鉴一起整理', 'blocks': [
('p', '新增道纹图鉴，把图标和说明按行列出来，收录目前可用的十三种天赋、十九种词条与分叉道纹。可以按分类定位，也可以输入物品等级，查看对应的词条数值区间。'),
('fig', '11_道纹图鉴天赋', '图鉴中的天赋列表 图标左侧展示 规则右侧说明'),
('p', '力量、速度、智力归基础属性；血量、灵力、防御、护盾、攻速、移速归普通属性。形态仍是数量、分裂、连锁、弧度和范围。元素先集中做金木水火土，提供额外伤害；冰、雷、时间、空间暂时退出新生成池。元素控制和灼烧等独立效果还没接入。'),
('fig', '12_普通属性与成长', '普通属性列表与当前物品等级的数值区间')
]},
{'heading': '六 分叉道纹专门负责连接', 'blocks': [
('p', '接口也做了取舍。常规属性道纹天然只出1或2个接口，目前概率是70%和30%；再单独加入没有属性、拥有3到6个接口的分叉道纹，专门负责传导和分流。分叉不能改造，接口掉落后固定，可以通过旋转调整方向。'),
('fig', '13_分叉与连接规则', '图鉴中的分叉道纹 品阶容量与通路规则'),
('p', '八个品阶继续保留：普通、优秀、杰出、稀有、完美、超凡、史诗、传说。一般属性道纹的品阶控制词条容量，五行道纹固定单条额外伤害，同项词条允许重复并相加。品阶和接口数量各有职责，不再一起增加。'),
('p', '这一轮也把成长、攻击、防御、抗性、暴击和经验接入同一套正式数值。现在游戏展示和实战都用这些公式，仍然需要后续试玩调整平衡。“已经接通数值”和“已经平衡好了”是两回事，我还得继续测。')
]},
{'heading': '七 想打多少级 直接在主页选', 'blocks': [
('p', '选图从点击打开，改成了主页右下角常驻。三张方形卡片中，青岚原可以选择，后两张保持封印。先把第一张地图做扎实，后面的世界再逐步展开。'),
('fig', '04_主页地图', '主页常驻地图卡与黑色半透明信息面板'),
('p', '地图等级改为1到100级的下拉选择。普通怪与地图同级，精英高1级，头目高3级，BOSS高5级。玩家自己决定挑战多少级，进入前再确认一次详细信息。'),
('fig', '05_地图等级下拉', '地图等级逐级选择 默认从1级开始')
]},
{'heading': '八 青岚原改成固定开放战场', 'blocks': [
('p', '进入确认会列清楚这次要去的地图、敌人等级和数量。这些等级会影响实际敌人属性，选择高等级地图就要承担更高的挑战。'),
('fig', '14_进入地图确认', '进入前确认地图等级与敌人配置'),
('p', '游戏重新回归纯2D，青岚原是80×80米的固定战场，玩家从中央出发。移动区域按地图中的空地描出来，树林、河流和岩壁会挡路；通行边界有青绿色发光线。主角会按八个方向切换朝向，动作连贯性还需要继续打磨。'),
('fig', '15_纯2D战场与紧凑HUD', '青岚原实战 通行边界与四角战斗信息')
]},
{'heading': '九 小怪 精英 头目和狼王的节奏', 'blocks': [
('p', '一局预留66只敌人：60只普通、4只精英、1只头目、1只BOSS。普通怪分五组递增登场，穿插精英和头目。普通、精英和头目从屏幕外的可走区域出现，再寻路追向玩家；受伤后显示血条。'),
('fig', '16_伤害来源与敌人血条', '敌人追击 受伤血条与普通 火伤害分项显示'),
('p', '非BOSS敌人的实际生命损伤达到80%后，青鬃狼王在地图最北端出现，再朝玩家移动。精英有突进，头目有范围震圈，狼王会切换攻击，半血后加速。我想让构筑有输出空间，也让走位一直有事情可做。'),
('fig', '19_青鬃狼王挑战', '青鬃狼王登场后的真实挑战画面')
]},
{'heading': '十 战斗信息少挡一点 掉落少漏一点', 'blocks': [
('p', '生命、护盾和经验放左上，敌人总剩余、场上、待刷新与各类数量放右上，小地图和阶段目标放左下。资源和详细构筑放进暂停页，按Esc就能查看，暂停期间敌人和刷新也会停止。PC隐藏移动摇杆。'),
('fig', '18_战斗暂停详情', '暂停查看资源 经验和当前通路参数'),
('p', '伤害飘字拆成普通与五行来源：普通白、金金、木绿、水蓝、火红、土黄。道纹、通货、灵石无视距离自动入库，吸附只是演出；拾取消息逐条滑入、停留再滑出，快速击杀时也能看到拿到了什么。已入库奖励在角色倒下或离场后保留。'),
('fig', '17_自动吸附拾取', '掉落自动吸附入库 每条通知各自显示')
]},
{'heading': '十一 灵石和宝盒补上收集入口', 'blocks': [
('p', '灵石现在作为独立掉落接入：普通敌人5、精英25、头目75、BOSS200。回主页可以开启三类宝盒，属性宝盒100灵石、功能宝盒500灵石、分叉宝盒1000灵石，每次获得一枚道纹。'),
('fig', '21_宝盒与抽取日志', '宝盒界面 展示价格 概率与真实抽取记录'),
('p', '抽取日志会记下消耗和结果，鼠标悬停能看这枚道纹的词条、品阶和接口。这样战斗获得的灵石，回到主页后就有了直接用途。当前概率是原型数值，还会根据实际收集节奏调整。'),
('fig', '26_宝盒结果悬停详情', '悬停抽取结果 查看这枚实际道纹的详情')
]},
{'heading': '十二 改造界面和通货继续精简', 'blocks': [
('p', '改造界面把通货名字、图标和数量放到一起，数量放大；右侧专门展示目标道纹和词条，介绍移到最下方。希望打开界面后，先能看清自己有什么、准备改哪一枚。'),
('fig', '22_通货改造界面', '通货改造界面 数量醒目 目标详情单独排列'),
('p', '接口现在固定，原来的通脉针和六通玉退出可用列表。目前保留11种通货，负责定向升阶、补充词条、单条洗练、全部洗练和随机升阶。常规道纹升阶保留接口；图中演示的是消耗启灵石，将普通移速道纹升成优秀。'),
('fig', '23_改造结果', '实际消耗1枚启灵石后的优秀移速道纹')
]},
{'heading': '十三 从打一轮 到回来再构筑', 'blocks': [
('p', '这一版的循环已经能跑下来：选天赋、连接道纹、选地图等级、战斗获得经验和掉落，再回主页改造、解锁格子、重新调整。这次录制使用当前进度，打的是正式一级地图，清完66只敌人并击败狼王，角色从2级升到5级。'),
('fig', '20_战斗结果', '狼王击败后可返回 已获得的战利品保留'),
('p', '回到画布，新的技能点可以继续开格，新增的道纹也可以接到原有通路里。截图中把木道纹接到火道纹后方，右侧五行伤害随之改变。我希望“这枚东西能放到哪里”这个想法，能自然变成下一次战斗的尝试。'),
('fig', '24_成长后重新构筑', '成长后接入木道纹 有效连接与伤害同步变化'),
('p', '这次也录了一段最新版介绍视频，展示从配置道纹、进入战场到回主页继续调整的过程。战斗保留原速，跳切省去重复赶路；宝盒部分使用同版本补拍，并在画面中标出。')
]},
{'heading': '问题与风险', 'blocks': [
('p', '功能越来越多，不代表体验就已经到位了。当前最明显的还是角色动作：八方向朝向已经接入，但步幅、身体起伏和动作衔接还不够自然。命中反馈和精英、BOSS的攻击提示也要继续打磨，至少得让玩家看清自己打中了什么、什么时候该躲。'),
('p', '数值已经使用同一套正式公式，但平衡还没做完。这次跑通的是一级地图，不能据此说1到100级都验证好了。高等级敌人强度、道纹成长、灵石收入与宝盒消耗之间是否匹配，还需要逐段试玩；如果只靠堆数值才能过图，构筑的选择就没意义了。'),
('p', '内容上目前只开放青岚原，后两张地图仍然封印。五行先提供伤害，控制、灼烧等元素效果，以及条件特性道纹和转化组合，都没有进入当前版本。接下来还要把新增规则控制在能解释、能验证的范围内，避免系统越加越多，玩家反而看不懂。'),
('section', '下期计划'),
('p', '第一件事，继续修角色移动和战斗反馈。优先统一步幅与身体起伏，补清楚精英和狼王的攻击预警，再检查伤害飘字与拾取消息在混战中是否仍然好读。'),
('p', '第二件事，围绕现有地图做成长测试。先检查低等级的升级、掉落和改造节奏，再抽查中高等级挑战，记录哪里变得过慢、哪里出现收益失衡，按结果调整正式数值。'),
('p', '第三件事，把条件特性与转化组合收敛成一个可以验证的样例。先明确激活条件、下游形态归属与冲突处理，再接入最小组合，确认玩家能看懂它为什么生效；不会把所有设想一口气塞进游戏。'),
('p', '下一期就跟着这些实际进度继续更新。欢迎聊聊你想搭怎样的道纹通路，也欢迎指出目前看着别扭的地方。第二期先记录到这里，我是烤肠，我们下一期见。')
]}
]

def set_font(style, name='Microsoft YaHei', size=11):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.font.underline = False
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ('asciiTheme', 'eastAsiaTheme', 'hAnsiTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + attr), None)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), name)
    for tag in ('spacing', 'kern', 'iCs'):
        node = rpr.find(qn('w:' + tag))
        if node is not None:
            rpr.remove(node)
    size_cs = rpr.find(qn('w:szCs'))
    if size_cs is None:
        size_cs = OxmlElement('w:szCs')
        rpr.append(size_cs)
    size_cs.set(qn('w:val'), str(size * 2))
    ppr = style._element.find(qn('w:pPr'))
    if ppr is not None:
        for tag in ('pBdr', 'numPr'):
            node = ppr.find(qn('w:' + tag))
            if node is not None:
                ppr.remove(node)

def main():
    OUT.mkdir(exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(.7)
    sec.top_margin = sec.bottom_margin = Inches(.62)
    set_font(doc.styles['Normal'], size=11)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.15
    normal.space_after = Pt(6)
    set_font(doc.styles['Title'], size=20)
    doc.styles['Title'].paragraph_format.space_after = Pt(6)
    set_font(doc.styles['Subtitle'], size=16)
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(12)
    set_font(doc.styles['Heading 1'], size=16)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
    set_font(doc.styles['Heading 2'], size=16)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(0)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(9)
    set_font(doc.styles['Caption'], size=10)
    doc.styles['Caption'].paragraph_format.space_after = Pt(8)
    doc.styles['Caption'].paragraph_format.space_before = Pt(2)
    doc.styles['Caption'].font.bold = False
    grid = sec._sectPr.find(qn('w:docGrid'))
    if grid is not None:
        sec._sectPr.remove(grid)
    doc.core_properties.author = '烤肠'
    doc.core_properties.title = TITLE
    doc.core_properties.subject = '第二期开发日志与当前版本实机截图'
    doc.add_paragraph('从参加聚光灯比赛到我为天帝镇压世间一切', 'Title')
    doc.add_paragraph('第二期开发日志', 'Subtitle')
    md = [f'# 《从参加聚光灯比赛到我为天帝镇压世间一切》第二期开发日志', '']
    figure = 0
    mapping = []
    for pi, page in enumerate(PAGES):
        if pi:
            style = 'Heading 2' if pi <= 13 else 'Heading 1'
            h = doc.add_paragraph(page['heading'], style)
            h.paragraph_format.page_break_before = True
            md += [('### ' if pi <= 13 else '## ') + page['heading'], '']
        for typ, *rest in page['blocks']:
            if typ == 'p':
                p = doc.add_paragraph(rest[0])
                p.paragraph_format.widow_control = True
                md += [rest[0], '']
            elif typ == 'section':
                doc.add_paragraph(rest[0], 'Heading 1')
                md += ['## ' + rest[0], '']
            else:
                stem, label = rest
                path = OUT / '截图' / (stem + '.png')
                if not path.exists():
                    raise FileNotFoundError(path)
                figure += 1
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.line_spacing = 1
                # 双图页适当留文字空间，保持原图比例；全分辨率图片独立保存。
                width = 6.85 if pi in (0, 6) else 5.15 if pi == 13 else 6.1
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(path), width=Inches(width))
                props = run._r.xpath('.//wp:docPr')
                if props:
                    props[0].set('descr', label)
                caption = f'图{figure} {label}'
                cp = doc.add_paragraph(caption, 'Caption')
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                md += [f'![{caption}](截图/{stem}.png)', '']
                mapping.append({'图号':figure, '正文位置':page['heading'] or '开篇', '文件': '截图/'+stem+'.png', '图注':label})
    for p in doc.paragraphs:
        node = OxmlElement('w:snapToGrid')
        node.set(qn('w:val'), '0')
        p._p.get_or_add_pPr().append(node)
    doc.save(OUT / '第二期开发日志.docx')
    md += ['最新版介绍视频：[观看视频](视频/第二期最新版介绍.mp4)', '']
    (OUT / '第二期开发日志.md').write_text('\n'.join(md), encoding='utf-8')
    (OUT / '截图插入位置.md').write_text('# 截图插入位置\n\nWord已嵌入全部截图。发帖时，可按下列位置手动上传原图，顺序与正文一致。\n\n' + '\n'.join(f"{x['图号']}. {x['正文位置']} → `{x['文件']}`：{x['图注']}" for x in mapping)+'\n', encoding='utf-8')
    (ROOT / '生成/开发日志二期采集-20261005-185151/正文结构.json').write_text(json.dumps(PAGES, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'docx':str(OUT/'第二期开发日志.docx'),'images':figure,'expected_pages':len(PAGES),'paragraph_characters':sum(len(r[1]) for p in PAGES for r in p['blocks'] if r[0]=='p')}, ensure_ascii=False))

if __name__ == '__main__':
    main()
